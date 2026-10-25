import json
import os
import tempfile
import unittest
from datetime import date, datetime

from flask import Flask
from openpyxl import load_workbook

import app as school_app
from powerbi_reporting import (
    DATASET_COLUMNS,
    build_csv_export,
    build_excel_export,
    build_powerbi_datasets
)


class PowerBIReportingTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_app = Flask("isolated-powerbi-reporting-tests")
        self.test_app.config.update(
            SECRET_KEY="stable-reporting-test-key",
            SQLALCHEMY_DATABASE_URI=(
                "sqlite:///"
                + os.path.join(self.temp_dir.name, "powerbi-reporting.db")
            ),
            SQLALCHEMY_TRACK_MODIFICATIONS=False
        )
        school_app.db.init_app(self.test_app)
        with self.test_app.app_context():
            school_app.db.create_all()
            self.admin = school_app.User(
                name="Private Admin Name",
                username="private-admin-login",
                password="not-used",
                role="admin",
                active=True
            )
            self.teacher_user = school_app.User(
                name="Private Teacher Name",
                username="private-teacher-login",
                password="not-used",
                role="teacher",
                active=True
            )
            self.other_teacher_user = school_app.User(
                name="Private Other Teacher Name",
                username="private-other-teacher-login",
                password="not-used",
                role="teacher",
                active=True
            )
            self.class_one = school_app.SchoolClass(
                standard=7,
                stream="",
                division="A",
                academic_year="2026-27",
                capacity=30,
                active=True
            )
            self.class_two = school_app.SchoolClass(
                standard=7,
                stream="",
                division="B",
                academic_year="2026-27",
                capacity=30,
                active=True
            )
            self.class_three = school_app.SchoolClass(
                standard=7,
                stream="",
                division="C",
                academic_year="2026-27",
                capacity=30,
                active=True
            )
            self.subject = school_app.Subject(
                name="Mathematics",
                code="MATH-7"
            )
            school_app.db.session.add_all([
                self.admin,
                self.teacher_user,
                self.other_teacher_user,
                self.class_one,
                self.class_two,
                self.class_three,
                self.subject
            ])
            school_app.db.session.flush()
            self.teacher = school_app.Teacher(
                employee_id="PRIVATE-EMPLOYEE-ID",
                user_id=self.teacher_user.id,
                phone="555-PRIVATE",
                email="private-teacher@example.test",
                qualification="Private qualification"
            )
            self.other_teacher = school_app.Teacher(
                employee_id="PRIVATE-EMPLOYEE-ID-2",
                user_id=self.other_teacher_user.id,
                phone="555-PRIVATE-2",
                email="private-teacher-2@example.test"
            )
            self.student = school_app.Student(
                admission_number="PRIVATE-ADMISSION-ID",
                student_id="PRIVATE-STUDENT-ID",
                name="Private Student Name",
                email="private-student@example.test",
                phone="555-STUDENT-PRIVATE",
                address="Private street address",
                current_standard=7,
                class_id=self.class_one.id,
                admission_date=date(2026, 9, 1),
                admission_status="Confirmed",
                admission_fee=12345,
                fee_status="Paid",
                active=True
            )
            self.staff = school_app.Staff(
                employee_id="PRIVATE-STAFF-ID",
                name="Private Staff Name",
                position="Librarian",
                phone="555-STAFF-PRIVATE",
                active=True
            )
            school_app.db.session.add_all([
                self.teacher,
                self.other_teacher,
                self.student,
                self.staff
            ])
            school_app.db.session.flush()
            self.exam = school_app.Exam(
                name="Term Assessment",
                exam_date=date(2026, 9, 15)
            )
            school_app.db.session.add(self.exam)
            school_app.db.session.flush()
            school_app.db.session.add_all([
                school_app.Attendance(
                    student_id=self.student.id,
                    teacher_id=self.teacher.id,
                    attendance_date=date(2026, 9, 2),
                    status="Present"
                ),
                school_app.Mark(
                    student_id=self.student.id,
                    teacher_id=self.teacher.id,
                    subject_id=self.subject.id,
                    exam_id=self.exam.id,
                    marks_obtained=18,
                    maximum_marks=20,
                    absent=False,
                    created_at=datetime(2026, 9, 15, 9, 0)
                ),
                school_app.TeacherAssignment(
                    teacher_id=self.teacher.id,
                    class_id=self.class_one.id,
                    subject_id=self.subject.id
                ),
                school_app.LeaveRequest(
                    requester_type="teacher",
                    teacher_id=self.teacher.id,
                    leave_type="Personal",
                    start_date=date(2026, 9, 20),
                    end_date=date(2026, 9, 21),
                    reason="Private leave reason",
                    review_note="Private reviewer note",
                    status="Approved"
                ),
                school_app.Timetable(
                    class_id=self.class_one.id,
                    teacher_id=self.teacher.id,
                    subject_id=self.subject.id,
                    day="Monday",
                    period_number=1,
                    start_time="09:00",
                    end_time="09:45",
                    room="Lab 1"
                ),
                school_app.Timetable(
                    class_id=self.class_two.id,
                    teacher_id=self.teacher.id,
                    subject_id=self.subject.id,
                    day="Monday",
                    period_number=2,
                    start_time="09:30",
                    end_time="10:00",
                    room="Lab 2"
                ),
                school_app.Timetable(
                    class_id=self.class_three.id,
                    teacher_id=self.other_teacher.id,
                    subject_id=self.subject.id,
                    day="Monday",
                    period_number=3,
                    start_time="09:15",
                    end_time="09:35",
                    room=" lab 1 "
                ),
                school_app.TimetableDraft(
                    class_id=self.class_one.id,
                    required_periods=2,
                    generated_periods=1,
                    preference_percentage=75,
                    quality_score=82,
                    status="Draft",
                    created_by=self.admin.id
                ),
                school_app.GoogleSheetSyncBatch(
                    id="private-sheet-batch-id",
                    record_type="admissions",
                    spreadsheet_id="private-spreadsheet-id",
                    sheet_range="Admissions!A1:I2",
                    status="Approved",
                    row_count=1,
                    applied_count=1,
                    created_by=self.admin.id
                ),
                school_app.GoogleSheetExport(
                    record_type="admissions",
                    record_id=self.student.id,
                    attempts=2,
                    last_error="Private integration error"
                ),
                school_app.KafkaOutboxEvent(
                    event_id="private-kafka-event-id",
                    event_type="student.attendance.recorded",
                    entity_id="private-entity-id",
                    payload_json='{"private":"payload"}',
                    attempts=1
                ),
                school_app.StudentModelTrainingJob(
                    job_id="private-training-job-id",
                    uploaded_by_id=self.admin.id,
                    original_filename="private-students.csv",
                    storage_filename="private-storage-key.csv",
                    dataset_sha256="private-dataset-hash",
                    status="Trained",
                    row_count=50,
                    metrics_json=json.dumps({
                        "metrics": {
                            "pass_fail": {
                                "accuracy": 0.8,
                                "balanced_accuracy": 0.75
                            },
                            "final_percentage": {
                                "mae_percentage_points": 4.2,
                                "r2": 0.61
                            }
                        },
                        "private_unapproved_field": "must not export"
                    })
                )
            ])
            school_app.db.session.commit()
            self.models = school_app._powerbi_reporting_models()

    def tearDown(self):
        with self.test_app.app_context():
            school_app.db.session.remove()
            school_app.db.drop_all()
            school_app.db.engine.dispose()
        self.temp_dir.cleanup()

    def datasets(self):
        return build_powerbi_datasets(
            self.models,
            self.test_app.config["SECRET_KEY"]
        )

    def test_star_schema_datasets_have_consistent_keys_and_accurate_measures(self):
        with self.test_app.app_context():
            datasets = self.datasets()

        self.assertEqual(
            list(datasets["fact_attendance"][0]),
            DATASET_COLUMNS["fact_attendance"]
        )
        student_key = datasets["dim_student"][0]["student_key"]
        self.assertEqual(
            datasets["fact_attendance"][0]["student_key"],
            student_key
        )
        self.assertEqual(
            datasets["fact_admissions"][0]["date_key"],
            20260901
        )
        self.assertEqual(
            datasets["dim_class"][0]["enrollment_count"],
            1
        )
        self.assertEqual(
            datasets["fact_exam_results"][0]["score_percent"],
            90
        )
        self.assertEqual(
            datasets["fact_ml_training"][0]["pass_fail_accuracy"],
            0.8
        )
        self.assertEqual(
            datasets["fact_teacher_availability"][0]["leave_days"],
            2
        )

    def test_sensitive_personal_and_integration_values_are_excluded(self):
        with self.test_app.app_context():
            datasets = self.datasets()
        serialized = json.dumps(datasets)
        for sensitive in (
            "Private Student Name",
            "PRIVATE-ADMISSION-ID",
            "PRIVATE-STUDENT-ID",
            "private-student@example.test",
            "Private Teacher Name",
            "PRIVATE-EMPLOYEE-ID",
            "private-teacher@example.test",
            "Private Other Teacher Name",
            "PRIVATE-EMPLOYEE-ID-2",
            "private-teacher-2@example.test",
            "Private Staff Name",
            "PRIVATE-STAFF-ID",
            "Private leave reason",
            "Private reviewer note",
            "private-spreadsheet-id",
            "private-sheet-batch-id",
            "private-kafka-event-id",
            "private-entity-id",
            "private-dataset-hash",
            "private-students.csv",
            "Private integration error",
            "must not export"
        ):
            self.assertNotIn(sensitive, serialized)

    def test_timetable_conflicts_are_detected_by_actual_overlapping_times(self):
        with self.test_app.app_context():
            conflicts = self.datasets()["fact_timetable_conflicts"]
        conflict_types = {row["conflict_type"] for row in conflicts}
        self.assertEqual(conflict_types, {"teacher_overlap", "room_overlap"})
        self.assertEqual(len(conflicts), 2)

    def test_operations_exports_are_aggregated_without_payloads_or_ids(self):
        with self.test_app.app_context():
            datasets = self.datasets()
        self.assertEqual(
            datasets["agg_google_sheets_sync"][0]["applied_count"],
            1
        )
        self.assertEqual(
            datasets["agg_google_sheets_exports"][0]["export_status"],
            "Retry scheduled"
        )
        self.assertEqual(
            datasets["agg_kafka_processing"][0]["processing_status"],
            "Pending"
        )

    def test_csv_and_excel_exports_are_usable_and_formula_safe(self):
        with self.test_app.app_context():
            datasets = self.datasets()
        csv_content = build_csv_export(
            "dim_subject",
            [{"subject_key": 1, "subject_code": "=1+1", "subject_name": "Math"}]
        )
        self.assertIn("'=1+1", csv_content)

        datasets["dim_subject"] = [{
            "subject_key": 1,
            "subject_code": "=1+1",
            "subject_name": "Mathematics"
        }]
        workbook_file = build_excel_export(datasets)
        workbook = load_workbook(workbook_file, read_only=False, data_only=False)
        self.assertIn("dim_student", workbook.sheetnames)
        self.assertEqual(
            workbook["dim_student"].max_row,
            len(datasets["dim_student"]) + 1
        )
        self.assertEqual(workbook["dim_subject"]["B2"].value, "'=1+1")
        self.assertEqual(workbook["dim_subject"]["B2"].data_type, "s")
        workbook.close()

    def test_powerbi_download_page_requires_login(self):
        client = school_app.app.test_client()
        for path in (
            "/admin/power-bi",
            "/admin/power-bi/export/dim_class.csv",
            "/admin/power-bi/export.xlsx"
        ):
            with self.subTest(path=path):
                response = client.get(path)
                self.assertEqual(response.status_code, 302)
                self.assertIn("/login", response.headers["Location"])


if __name__ == "__main__":
    unittest.main()
